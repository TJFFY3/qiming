import copy
import io
import json
from types import SimpleNamespace

import pytest
from docx import Document
from fastapi.testclient import TestClient

PREFIX = "/api/timed-practice"


def paper(title="氧化还原反应"):
    return {
        "title": title,
        "duration": 20,
        "total_score": 50,
        "choice": [
            dict(
                stem=f"判断第{i}个反应中元素化合价变化。",
                options=["甲", "乙", "丙", "丁"],
                answer="A",
                explanation="比较反应前后化合价。",
                score=4,
                minutes=1,
                difficulty="基础",
                source_ids=["s1"],
                origin="adapted",
            )
            for i in range(5)
        ],
        "written": [
            dict(
                stem=f"写出第{i}个反应的电子转移过程：________",
                answer="失电子，化合价升高。",
                explanation="依据化合价变化判断。",
                score=15,
                minutes=5,
                difficulty="中档",
                source_ids=["s1"],
                origin="adapted",
                work_lines=3,
            )
            for i in range(2)
        ],
        "optional": [],
        "source_gaps": [],
        "conflicts": [],
    }


def source(verified=False):
    return dict(
        id="s1",
        category="textbook",
        title="教材",
        locator="检索片段",
        evidence="比较元素反应前后的化合价，分析电子转移。",
        verified=verified,
    )


def settings():
    from backend.app.timed_practice.models import LessonSettings

    return LessonSettings(
        options={
            "duration": "20 分钟",
            "score": "50 分",
            "questions": "5 道选择题和 2 道解答题",
        }
    ).model_dump()


def test_custom_contract_and_two_separate_word_files(tmp_path):
    from backend.app.timed_practice.requirements import resolve_requirements
    from backend.app.timed_practice.documents import validate_content, render

    constraints = resolve_requirements(settings(), "请生成限时练习")
    content = validate_content(paper(), constraints, [source()])
    report = render(content, [source()], tmp_path / "v1", True)
    student = Document(tmp_path / "v1" / "paper.docx")
    answers = Document(tmp_path / "v1" / "answers.docx")
    assert len(student.tables[0].columns) == 6
    assert "比较反应前后化合价" not in "\n".join(p.text for p in student.paragraphs)
    assert "比较反应前后化合价" in "\n".join(p.text for p in answers.paragraphs)
    assert "50分" in "\n".join(p.text for p in answers.paragraphs)
    assert "待核验" not in "\n".join(
        p.text for p in student.paragraphs + answers.paragraphs
    )
    assert "待核验" not in json.dumps(report, ensure_ascii=False)
    assert report["draft"] is False
    assert report["validation"]["passed"] is True
    assert (
        json.loads((tmp_path / "v1" / "content.json").read_text())["total_score"] == 50
    )


def test_invalid_score_sources_origin_and_paths_rejected():
    from backend.app.timed_practice.requirements import resolve_requirements
    from backend.app.timed_practice.documents import validate_content

    c = resolve_requirements(settings(), "")
    for key, value in [
        ("score", 3),
        ("source_ids", ["fake"]),
        ("image", "/etc/passwd"),
        ("stem", ""),
        ("minutes", float("nan")),
    ]:
        bad = paper()
        bad["choice"][0][key] = value
        with pytest.raises(ValueError):
            validate_content(bad, c, [source()])
    original = dict(c, origin="original")
    with pytest.raises(ValueError, match="原题"):
        validate_content(paper(), original, [source()])
    bad = paper()
    bad["choice"][0]["stem"] = [{"text": "x", "sub": True, "sup": True}]
    with pytest.raises(ValueError):
        validate_content(bad, c, [source()])


def test_requirements_explicit_body_overrides_options_and_default_profile():
    from backend.app.timed_practice.requirements import resolve_requirements
    from backend.app.timed_practice.models import LessonSettings

    c = resolve_requirements(settings(), "改为30分钟，满分100分，3道选择题和2道解答题")
    assert (c["duration"], c["total_score"], c["choice_count"], c["written_count"]) == (
        30,
        100,
        3,
        2,
    )
    d = resolve_requirements(LessonSettings().model_dump(), "")
    assert (
        d["duration"],
        d["total_score"],
        d["choice_count"],
        d["optional_count"],
    ) == (35, 100, 14, 2)
    inherited = resolve_requirements(LessonSettings().model_dump(), "", paper())
    assert (inherited["duration"], inherited["total_score"]) == (20, 50)


def app_for(tmp_path, replies=None):
    from backend.app.config import Settings
    from backend.app.main import create_app
    from backend.app.timed_practice.pipeline import PracticePipeline

    class Model:
        def generate_json(self, messages):
            return copy.deepcopy((replies or [paper()]).pop(0))

    def factory(config, store):
        return PracticePipeline(
            config, store, SimpleNamespace(search=lambda q: [source()]), Model()
        )

    return create_app(Settings(data_dir=tmp_path), practice_factory=factory)


def test_independent_namespace_unverified_generation_download_and_revision(tmp_path):
    app = app_for(tmp_path, [paper(), paper("修改后的练习")])
    with TestClient(app) as client:
        guide = client.post("/api/sessions", json={}).json()["id"]
        s = client.post(PREFIX + "/sessions", json=settings()).json()
        sid = s["id"]
        assert client.get(PREFIX + "/sessions/" + guide).status_code == 404
        assert client.get("/api/sessions/" + sid).status_code == 404
        result = client.post(
            f"{PREFIX}/sessions/{sid}/generate",
            json={"prompt": "生成练习", "allow_draft": False},
        ).json()
        job = client.get(PREFIX + "/jobs/" + result["id"]).json()
        assert job["status"] == "completed", job
        s = client.get(f"{PREFIX}/sessions/{sid}").json()
        v = s["versions"][0]
        assert "directory" not in v
        assert v["sources"][0]["verified"] is False
        assert v["report"]["draft"] is False
        assert "待核验" not in json.dumps(v, ensure_ascii=False)
        for kind in ["paper", "answers"]:
            response = client.get(
                f"{PREFIX}/sessions/{sid}/versions/{v['id']}/files/{kind}"
            )
            assert response.status_code == 200
            assert "待核验" not in response.headers["content-disposition"]
            assert "待核验" not in "\n".join(
                p.text for p in Document(io.BytesIO(response.content)).paragraphs
            )
        assert (
            client.get(
                f"{PREFIX}/sessions/{guide}/versions/{v['id']}/files/paper"
            ).status_code
            == 404
        )
        revision = client.post(
            f"{PREFIX}/sessions/{sid}/revise",
            json={
                "prompt": "修改标题",
                "base_version_id": v["id"],
                "allow_draft": True,
            },
        ).json()
        assert (
            client.get(PREFIX + "/jobs/" + revision["id"]).json()["status"]
            == "completed"
        )
        s2 = client.get(f"{PREFIX}/sessions/{sid}").json()
        assert s2["versions"][0] == v
        assert s2["versions"][1]["parent_id"] == v["id"]
        assert (
            client.post(
                f"{PREFIX}/sessions/{sid}/revise",
                json={"prompt": "再次修改", "base_version_id": v["id"]},
            ).status_code
            == 409
        )
    assert (tmp_path / "timed-practice" / "practice.sqlite").is_file()


def test_upload_delete_atomic_reservation_and_restart(tmp_path):
    from backend.app.timed_practice.store import PracticeStore, PracticeConflict

    store = PracticeStore(tmp_path / "independent.sqlite")
    s = store.create_session(settings())
    j = store.create_job(s["id"], {"prompt": "生成", "kind": "generation"})
    with pytest.raises(PracticeConflict):
        store.create_job(s["id"], {"prompt": "生成", "kind": "generation"})
    with pytest.raises(PracticeConflict):
        store.change_idle(s["id"], lambda s: s.update(resources=[]))
    assert store.recover_jobs() == 1
    assert store.get_job(j["id"])["status"] == "failed"
    with TestClient(app_for(tmp_path)) as client:
        sid = client.post(PREFIX + "/sessions", json={}).json()["id"]
        uploaded = client.post(
            f"{PREFIX}/sessions/{sid}/resources",
            files={"file": ("教材.txt", "原始资料".encode())},
            data={"category": "textbook"},
        ).json()
        assert "file_path" not in uploaded
        assert (
            client.post(
                f"{PREFIX}/sessions/{sid}/sources/{uploaded['id']}/review",
                json={"verified": True, "note": "已核对原文"},
            ).status_code
            == 200
        )
        assert (
            client.delete(
                f"{PREFIX}/sessions/{sid}/resources/{uploaded['id']}"
            ).status_code
            == 200
        )
        assert client.get(f"{PREFIX}/sessions/{sid}").json()["resources"] == []


def test_requirement_precedence_and_multiple_written_types():
    from backend.app.timed_practice.requirements import resolve_requirements

    s = settings()
    s["options"]["origin"] = "允许改编和出新题"
    c = resolve_requirements(s, "只用原题，2道填空题和1道解答题")
    assert c["origin"] == "original"
    assert c["written_count"] == 3
    s["options"]["extra"] = "用时30分钟，满分100分"
    c = resolve_requirements(s, "")
    assert (c["duration"], c["total_score"]) == (30, 100)


def test_default_time_profile_and_missing_optional_options_rejected():
    from backend.app.timed_practice.documents import validate_content
    from backend.app.timed_practice.requirements import resolve_requirements
    from backend.app.timed_practice.models import LessonSettings

    d = paper()
    d["duration"] = 35
    d["total_score"] = 100
    d["choice"] = [dict(d["choice"][0], stem=f"第{i}题", score=4) for i in range(14)]
    d["written"] = [dict(d["written"][0], score=44, minutes=1)]
    d["optional"] = [dict(d["written"][0], stem=f"选做{i}", score=0) for i in range(2)]
    with pytest.raises(ValueError, match="30"):
        validate_content(
            d, resolve_requirements(LessonSettings().model_dump(), ""), [source()]
        )


def test_repair_limit_and_noop_revision_never_publish_files(tmp_path):
    bad = paper()
    bad["choice"][0]["score"] = 1
    app = app_for(tmp_path, [bad, bad, bad])
    with TestClient(app) as client:
        sid = client.post(PREFIX + "/sessions", json=settings()).json()["id"]
        j = client.post(
            f"{PREFIX}/sessions/{sid}/generate",
            json={"prompt": "生成", "allow_draft": True},
        ).json()
        result = client.get(PREFIX + "/jobs/" + j["id"]).json()
        assert result["status"] == "failed" and "两轮" in result["error"]
        assert client.get(f"{PREFIX}/sessions/{sid}").json()["versions"] == []
        assert not list((tmp_path / "timed-practice" / sid).glob("versions/*/*.docx"))


def test_reviewed_requirements_survive_revision_and_editable_block_is_authoritative():
    from backend.app.timed_practice.requirements import resolve_requirements
    from backend.app.timed_practice.models import LessonSettings

    defaults = LessonSettings().model_dump()
    original = resolve_requirements(
        defaults, "20分钟，满分50分，5道选择题和2道解答题，只用原题"
    )
    revised = resolve_requirements(defaults, "修改第1题解析", paper(), original)
    assert (revised["origin"], revised["choice_count"], revised["written_count"]) == (
        "original",
        5,
        2,
    )
    prompt = "生成练习\n【限时练习优化设置】\n用时：30 分钟\n满分：100 分\n题型和数量：3道选择题和2道解答题\n出题方式：只用原题\n【限时练习优化设置结束】"
    result = resolve_requirements(settings(), prompt)
    assert (
        result["duration"],
        result["total_score"],
        result["choice_count"],
        result["origin"],
    ) == (30, 100, 3, "original")
    replacement = resolve_requirements(
        defaults,
        "原来14道选择题，改为5道选择题和2道填空题、1道解答题，20分钟，满分50分",
    )
    assert (replacement["choice_count"], replacement["written_count"]) == (5, 3)


def test_revision_uses_new_resources_and_current_reviews_without_mutating_old_version(
    tmp_path,
):
    from backend.app.timed_practice.pipeline import PracticePipeline
    from backend.app.timed_practice.store import PracticeStore
    from backend.app.config import Settings

    store = PracticeStore(tmp_path / "practice.sqlite")
    s = store.create_session(settings())
    captured = []

    class Model:
        def generate_json(self, messages):
            captured.append(json.loads(messages[1]["content"]))
            return paper() if len(captured) == 1 else paper("使用新资料的版本")

    config = Settings(data_dir=tmp_path)
    pipeline = PracticePipeline(
        config, store, SimpleNamespace(search=lambda q: [source()]), Model()
    )
    j = store.create_job(
        s["id"], dict(prompt="生成", kind="generation", allow_draft=True)
    )
    pipeline.run(j["id"])
    base = store.get_session(s["id"])["versions"][0]

    def update(current):
        current["sources"][0]["verified"] = True
        current["resources"].append(
            dict(
                source(True),
                id="new-upload",
                evidence="新的教师用书原文",
                category="teacher_book",
            )
        )

    store.change_idle(s["id"], update)
    j = store.create_job(
        s["id"],
        dict(
            prompt="使用新资料修改解析",
            kind="revision",
            base_version_id=base["id"],
            allow_draft=True,
        ),
    )
    pipeline.run(j["id"])
    assert store.get_job(j["id"])["status"] == "completed"
    sources = {x["id"]: x for x in captured[1]["sources"]}
    assert (
        sources["s1"]["verified"] is True
        and sources["new-upload"]["evidence"] == "新的教师用书原文"
    )
    assert store.get_session(s["id"])["versions"][0] == base


def test_interrupted_worker_cannot_publish_a_completed_version(tmp_path):
    from backend.app.timed_practice.pipeline import PracticePipeline
    from backend.app.timed_practice.store import PracticeStore
    from backend.app.config import Settings

    store = PracticeStore(tmp_path / "practice.sqlite")
    s = store.create_session(settings())

    class Model:
        def generate_json(self, messages):
            store.recover_jobs()
            return paper()

    pipeline = PracticePipeline(
        Settings(data_dir=tmp_path),
        store,
        SimpleNamespace(search=lambda q: [source()]),
        Model(),
    )
    j = store.create_job(
        s["id"], dict(prompt="生成", kind="generation", allow_draft=True)
    )
    pipeline.run(j["id"])
    assert store.get_job(j["id"])["status"] == "failed"
    assert store.get_session(s["id"])["versions"] == []


def test_public_job_does_not_offer_draft_retries():
    from backend.app.timed_practice.api import public_job

    job = public_job(
        {"id": "j", "request": {"kind": "generation", "allow_draft": True}}
    )
    assert "allow_draft" not in job
    assert "request" not in job


def test_realistic_model_option_labels_and_legacy_conditions_are_normalized(tmp_path):
    from backend.app.timed_practice.documents import validate_content, render
    from backend.app.timed_practice.requirements import resolve_requirements

    raw = paper()
    raw["choice"][0]["options"][0] = "A. CaCO₃ ===高温=== CaO + CO₂↑"
    content = validate_content(raw, resolve_requirements(settings(), ""), [source()])
    option = content["choice"][0]["options"][0]
    assert isinstance(option, list) and any(r.get("condition") for r in option)
    assert not "".join(r["text"] for r in option).startswith("A.")
    render(content, [source()], tmp_path / "normalized", True)
    xml = (tmp_path / "normalized" / "paper.docx").read_bytes()
    import zipfile

    with zipfile.ZipFile(io.BytesIO(xml)) as package:
        document = package.read("word/document.xml").decode()
        assert "<m:limUpp>" in document
        assert "<m:t>高温</m:t>" in document
        assert "===高温===" not in document


def test_invalid_json_retries_without_retrying_transport_failures(tmp_path):
    from backend.app.timed_practice.model_client import ModelOutputError
    from backend.app.timed_practice.pipeline import PracticePipeline
    from backend.app.timed_practice.store import PracticeStore
    from backend.app.config import Settings

    store = PracticeStore(tmp_path / "practice.sqlite")
    s = store.create_session(settings())

    class Model:
        calls = 0

        def generate_json(self, messages):
            self.calls += 1
            if self.calls == 1:
                raise ModelOutputError("模型返回无效 JSON")
            return paper()

    model = Model()
    pipeline = PracticePipeline(
        Settings(data_dir=tmp_path),
        store,
        SimpleNamespace(search=lambda q: [source()]),
        model,
    )
    j = store.create_job(
        s["id"], dict(prompt="生成", kind="generation", allow_draft=True)
    )
    pipeline.run(j["id"])
    assert store.get_job(j["id"])["status"] == "completed" and model.calls == 2


def test_model_network_failure_is_not_automatically_retried(tmp_path):
    from backend.app.timed_practice.pipeline import PracticePipeline
    from backend.app.timed_practice.store import PracticeStore
    from backend.app.config import Settings

    store = PracticeStore(tmp_path / "practice.sqlite")
    s = store.create_session(settings())

    class Model:
        calls = 0

        def generate_json(self, messages):
            self.calls += 1
            raise ValueError("模型响应超时，请重试")

    model = Model()
    pipeline = PracticePipeline(
        Settings(data_dir=tmp_path),
        store,
        SimpleNamespace(search=lambda q: [source()]),
        model,
    )
    j = store.create_job(
        s["id"], dict(prompt="生成", kind="generation", allow_draft=True)
    )
    pipeline.run(j["id"])
    assert store.get_job(j["id"])["status"] == "failed" and model.calls == 1


def test_semantic_run_defaults_do_not_erase_subscripts_and_normalization_is_idempotent():
    from backend.app.timed_practice.documents import validate_content
    from backend.app.timed_practice.requirements import resolve_requirements

    raw = paper()
    raw["written"][0]["stem"] = [{"text": "2[H2O] 与 [SO4^2-] 的变化：________"}]
    raw["choice"][0]["options"][0] = "A. H₂O 与 Fe³⁺"
    requirements = resolve_requirements(settings(), "")
    result = validate_content(raw, requirements, [source()])
    runs = result["written"][0]["stem"]
    assert any(r.get("sub") and r["text"] == "2" for r in runs)
    assert any(r.get("charge") and r["text"] == "2-" for r in runs)
    assert isinstance(result["choice"][0]["options"][0], list)
    assert validate_content(result, requirements, [source()]) == result


def test_unlocatable_original_is_never_published_as_original_when_new_questions_are_allowed():
    from backend.app.timed_practice.documents import validate_content, source_issues
    from backend.app.timed_practice.requirements import resolve_requirements

    raw = paper()
    raw["choice"][0]["origin"] = "original"
    requirements = resolve_requirements(settings(), "允许改编和出新题")
    result = validate_content(raw, requirements, [source()])
    assert result["choice"][0]["origin"] == "new"
    assert any("未能定位完整原题" in gap for gap in result["source_gaps"])
    assert any("自编" in issue for issue in source_issues(result, [source()]))
    assert validate_content(result, requirements, [source()]) == result
    with pytest.raises(ValueError, match="原题"):
        validate_content(raw, {**requirements, "origin": "original"}, [source()])
    with pytest.raises(ValueError, match="原题"):
        validate_content(raw, {**requirements, "origin": "adapted"}, [source()])


def test_template_and_format_samples_cannot_be_claimed_as_question_sources():
    from backend.app.timed_practice.documents import validate_content
    from backend.app.timed_practice.requirements import resolve_requirements

    template = dict(source(), title="限时练模板.docx")
    requirements = resolve_requirements(settings(), "允许改编和出新题")
    result = validate_content(paper(), requirements, [template])
    assert all(
        q["origin"] == "new" and q["source_ids"] == []
        for q in result["choice"] + result["written"]
    )
    assert any("模板" in gap for gap in result["source_gaps"])
    with pytest.raises(ValueError, match="题源"):
        validate_content(paper(), {**requirements, "origin": "adapted"}, [template])


def test_no_sources_still_generates_when_new_questions_are_allowed(tmp_path):
    from backend.app.config import Settings
    from backend.app.timed_practice.pipeline import PracticePipeline
    from backend.app.timed_practice.store import PracticeStore

    store = PracticeStore(tmp_path / "practice.sqlite")
    lesson = settings()
    lesson["options"]["origin"] = "允许改编和出新题"
    session = store.create_session(lesson)
    content = paper()
    for question in content["choice"] + content["written"]:
        question.update(source_ids=[], origin="new")
    pipeline = PracticePipeline(
        Settings(data_dir=tmp_path),
        store,
        SimpleNamespace(search=lambda query: []),
        SimpleNamespace(generate_json=lambda messages: content),
    )
    job = store.create_job(
        session["id"], dict(prompt="生成", kind="generation", allow_draft=False)
    )
    pipeline.run(job["id"])
    assert store.get_job(job["id"])["status"] == "completed"
    version = store.get_session(session["id"])["versions"][0]
    assert version["sources"] == []
    assert version["report"]["draft"] is False
    assert "待核验" not in json.dumps(version, ensure_ascii=False)


@pytest.mark.parametrize("verified", [False, True])
def test_retrieved_source_verification_is_preserved(tmp_path, verified):
    from backend.app.config import Settings
    from backend.app.timed_practice.pipeline import PracticePipeline
    from backend.app.timed_practice.store import PracticeStore

    store = PracticeStore(tmp_path / "practice.sqlite")
    session = store.create_session(settings())
    pipeline = PracticePipeline(
        Settings(data_dir=tmp_path),
        store,
        SimpleNamespace(search=lambda query: [source(verified)]),
        SimpleNamespace(generate_json=lambda messages: paper()),
    )
    job = store.create_job(
        session["id"], dict(prompt="生成", kind="generation", allow_draft=False)
    )
    pipeline.run(job["id"])
    assert store.get_job(job["id"])["status"] == "completed"
    assert (
        store.get_session(session["id"])["versions"][0]["sources"][0]["verified"]
        is verified
    )
