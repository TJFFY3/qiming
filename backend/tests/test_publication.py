from zipfile import ZipFile

from docx import Document


def test_historical_office_download_removes_split_marker_without_changing_original(
    tmp_path,
):
    from backend.app.shared.tools.publication import clean_office_download

    original = tmp_path / "old.docx"
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("待核")
    p.add_run("验 · ")
    p.add_run("氧化还原反应 H")
    p.add_run("2").font.subscript = True
    p.add_run("O")
    doc.save(original)
    before = original.read_bytes()
    public = clean_office_download(original)
    assert original.read_bytes() == before
    output = Document(public)
    assert output.paragraphs[0].text == "氧化还原反应 H2O"
    assert any(r.text == "2" and r.font.subscript for r in output.paragraphs[0].runs)
    assert "待核验" not in output.paragraphs[0].text
    assert clean_office_download(original) == public


def test_public_cleaning_removes_labels_but_preserves_evidence_and_requests():
    from backend.app.shared.tools.publication import clean_public

    value = {
        "content": {"title": "待核验·氧化还原反应"},
        "report": {
            "draft": True,
            "status": "draft_missing_sources",
            "source_issues": [
                "未提供已核对的教材或教师用书（verified=true）",
                "缺少实验图",
            ],
            "visual_review": "Office/WPS效果待核验",
        },
        "sources": [{"verified": False, "evidence": "原文包含待核验一词"}],
        "messages": [{"role": "user", "content": "取消待核验标识"}],
    }
    result = clean_public(value)
    assert result["content"]["title"] == "氧化还原反应"
    assert result["report"]["source_issues"] == ["缺少实验图"]
    assert "待核验" not in result["report"].get("visual_review", "")
    assert result["sources"] == value["sources"]
    assert result["messages"] == value["messages"]
    assert value["content"]["title"] == "待核验·氧化还原反应"


def test_historical_pptx_download_cleans_cover_and_preserves_slide_order(tmp_path):
    import sys
    from pathlib import Path

    from backend.app.shared.tools.publication import clean_office_download

    sys.path.insert(0, str(Path("vendor/chemistry/scripts").resolve()))
    from ppt_tools import generate_pptx

    path = tmp_path / "old.pptx"
    generate_pptx(
        [
            {"title": "待核验·反应", "stage": "启", "paragraphs": ["2[H2O]"]},
            {"title": "电子转移", "stage": "析", "paragraphs": ["[Fe^3+]"]},
        ],
        path,
    )
    original = path.read_bytes()
    with ZipFile(clean_office_download(path)) as z:
        cover = z.read("ppt/slides/slide1.xml").decode()
        second = z.read("ppt/slides/slide2.xml").decode()
        assert "待核验" not in cover
        assert "反应" in cover
        assert 'baseline="-25000"' in cover
        assert "电子转移" in second
        assert 'baseline="30000"' in second
    assert path.read_bytes() == original


def test_public_gap_notes_keep_actual_missing_material_while_removing_status():
    from backend.app.shared.tools.publication import clean_public

    value = {
        "source_gaps": [
            "第1题未能定位完整原题，按自编待核验题处理；引用仅作知识点参考。"
        ],
        "source_issues": ["缺少实验图，待核验", "教材页码尚待核对"],
    }
    public = clean_public(value)
    assert public["source_gaps"] == [
        "第1题未能定位完整原题，按自编题处理；引用仅作知识点参考。"
    ]
    assert any("缺少实验图" in note for note in public["source_issues"])
    assert any("教材页码" in note for note in public["source_issues"])
    assert "待核验" not in str(public)
    assert "待核对" not in str(public)


def test_repeated_and_unbracketed_trailing_review_markers_are_removed(tmp_path):
    from backend.app.shared.tools.publication import clean_title, clean_office_download

    assert clean_title("待核验 · 待核验-氧化还原反应 · 待核验") == "氧化还原反应"
    original = tmp_path / "legacy.docx"
    doc = Document()
    p = doc.add_paragraph()
    for value in ("待核", "验 · 待核验-", "氧化还原反应", " · 待核验"):
        p.add_run(value)
    doc.save(original)
    before = original.read_bytes()
    public = Document(clean_office_download(original))
    assert public.paragraphs[0].text == "氧化还原反应"
    assert original.read_bytes() == before
