import copy
import json
from pathlib import Path
import pytest
from backend.app.config import Settings
from backend.app.study_guide.storage.store import Store, Conflict
from backend.app.study_guide.application.pipeline import Pipeline
from backend.app.study_guide.tools.documents import DocumentEngine

SAMPLE = json.loads(Path("vendor/chemistry/examples/outline.json").read_text())


class Retriever:
    def search(self, query):
        return [
            {
                "id": "src",
                "title": "教材",
                "category": "textbook",
                "evidence": "元素化合价变化",
                "locator": "切片1",
                "verified": False,
            }
        ]


class Model:
    def __init__(self):
        self.count = 0

    def generate_json(self, messages):
        self.count += 1
        return copy.deepcopy(SAMPLE)


def setup(tmp_path, model=None):
    s = Settings(data_dir=tmp_path)
    store = Store(tmp_path / "db")
    session = store.create_session({"chapter": "氧化还原反应", "period": "第一课时"})
    return (
        store,
        session,
        Pipeline(s, store, Retriever(), model or Model(), DocumentEngine()),
    )


@pytest.mark.parametrize("allow_draft", [False, True])
def test_unverified_sources_do_not_block_generation(tmp_path, allow_draft):
    store, s, p = setup(tmp_path)
    job = p.start_generation(s["id"], {"prompt": "生成", "allow_draft": allow_draft})
    p.run(job["id"])
    done = store.get_job(job["id"])
    assert done["status"] == "completed", done
    version = store.get_session(s["id"])["versions"][-1]
    assert Path(version["directory"], "outline.docx").is_file()
    # Removing the gate must not fabricate verification of the underlying evidence.
    assert version["report"]["draft"]
    assert version["sources"][0]["verified"] is False


def test_source_gaps_and_conflicts_are_recorded_without_blocking_export(tmp_path):
    class WithGaps(Model):
        def generate_json(self, messages):
            value = super().generate_json(messages)
            value["source_gaps"] = ["教材实验细节未覆盖"]
            value["source_conflicts"] = ["参考资料描述有差异"]
            return value

    store, s, p = setup(tmp_path, WithGaps())
    job = p.start_generation(s["id"], {"prompt": "生成"})
    p.run(job["id"])
    done = store.get_job(job["id"])
    assert done["status"] == "completed", done
    version = store.get_session(s["id"])["versions"][-1]
    assert Path(version["directory"], "outline.docx").is_file()
    assert any(
        "教材实验细节未覆盖" in issue for issue in version["report"]["source_issues"]
    )


def test_repair_limit(tmp_path):
    class Bad(Model):
        def generate_json(self, messages):
            self.count += 1
            return {"outline": {"title": "bad"}}

    model = Bad()
    store, s, p = setup(tmp_path, model)
    job = p.start_generation(s["id"], {"prompt": "生成", "allow_draft": True})
    p.run(job["id"])
    assert model.count == 3
    assert store.get_job(job["id"])["status"] == "failed"
    assert not store.get_session(s["id"])["versions"]


def test_revision_and_stale_guard(tmp_path):
    model = Model()
    store, s, p = setup(tmp_path, model)
    j = p.start_generation(s["id"], {"prompt": "生成", "allow_draft": True})
    p.run(j["id"])
    v = store.get_session(s["id"])["versions"][-1]
    model.generate_json = lambda _: {
        "mode": "patch",
        "changes": [
            {
                "block_id": "question-1",
                "replacement": [{"type": "paragraph", "content": "请增加判断步骤。"}],
            }
        ],
    }
    j = p.start_revision(s["id"], {"prompt": "增加提示", "base_version_id": v["id"]})
    p.run(j["id"])
    assert store.get_job(j["id"])["status"] == "completed"
    latest = store.get_session(s["id"])["versions"][-1]
    assert latest["content"]["sections"][1:] == v["content"]["sections"][1:]
    with pytest.raises(Conflict):
        p.start_revision(s["id"], {"prompt": "stale", "base_version_id": v["id"]})


def test_full_revision_inherits_draft_even_if_title_marker_removed(tmp_path):
    model = Model()
    store, s, p = setup(tmp_path, model)
    raw = copy.deepcopy(SAMPLE)
    raw["source_gaps"] = []
    raw["outline"]["title"] = "待核验 原课时"
    source = Retriever().search("")[0]
    source["verified"] = True
    p.retriever.search = lambda _: [source]
    model.generate_json = lambda _: copy.deepcopy(raw)
    j = p.start_generation(s["id"], {"prompt": "生成", "allow_draft": True})
    p.run(j["id"])
    base = store.get_session(s["id"])["versions"][-1]
    assert base["report"]["draft"]
    changed = copy.deepcopy(raw)
    changed["mode"] = "full"
    changed["outline"]["title"] = "新课时"
    model.generate_json = lambda _: changed
    j = p.start_revision(s["id"], {"prompt": "更改标题", "base_version_id": base["id"]})
    p.run(j["id"])
    latest = store.get_session(s["id"])["versions"][-1]
    assert latest["id"] != base["id"]
    assert latest["report"]["draft"] is True


def test_completion_does_not_delete_published_artifacts(tmp_path):
    store, s, p = setup(tmp_path)
    original = store.update_job

    def fail_after_publish(job_id, **values):
        if store.get_job(job_id)["status"] == "completed":
            raise OSError("simulate DB failure")
        return original(job_id, **values)

    store.update_job = fail_after_publish
    j = p.start_generation(s["id"], {"prompt": "生成", "allow_draft": True})
    try:
        p.run(j["id"])
    except OSError:
        pass
    versions = store.get_session(s["id"])["versions"]
    assert len(versions) == 1
    assert Path(versions[0]["directory"], "outline.docx").exists()
    assert all(
        step["status"] == "completed" for step in store.get_job(j["id"])["stages"]
    )


def test_noop_revision_retries_and_does_not_publish_version(tmp_path):
    model = Model()
    store, s, p = setup(tmp_path, model)
    j = p.start_generation(s["id"], {"prompt": "生成", "allow_draft": True})
    p.run(j["id"])
    base = store.get_session(s["id"])["versions"][-1]
    first = base["content"]["sections"][0]["blocks"][0]
    calls = []

    def no_change(messages):
        calls.append(messages)
        return {
            "mode": "patch",
            "changes": [{"block_id": first["id"], "replacement": [first]}],
        }

    model.generate_json = no_change
    j = p.start_revision(s["id"], {"prompt": "增加提示", "base_version_id": base["id"]})
    p.run(j["id"])
    assert store.get_job(j["id"])["status"] == "failed"
    assert len(store.get_session(s["id"])["versions"]) == 1
    assert len(calls) == 3
    assert "实际" in store.get_job(j["id"])["error"]


def test_optimization_is_visible_in_stored_conversation_and_model_prompt(tmp_path):
    store, session, pipeline = setup(tmp_path)
    options = {
        "scope": "仅第一课时",
        "duration": "15 分钟",
        "tasks": "3 个任务",
        "difficulty": "基础为主",
        "hints": "详细提示",
        "focus": ["概念理解", "化学方程式"],
        "extra": "学生基础较弱",
    }
    store.update_session(session["id"], lambda s: s["settings"].update(options=options))
    captured = []

    def model(messages):
        captured.append(messages)
        return copy.deepcopy(SAMPLE)

    pipeline.model.generate_json = model
    job = pipeline.start_generation(session["id"], {"prompt": "以生活情境导入"})
    pipeline.run(job["id"])
    assert store.get_job(job["id"])["status"] == "completed"
    text = store.get_session(session["id"])["messages"][0]["content"]
    for line in [
        "学习范围：仅第一课时",
        "自学用时：15 分钟",
        "任务数量：3 个任务",
        "内容难度：基础为主",
        "提示详略：详细提示",
        "本课重点：概念理解、化学方程式",
        "补充要求：学生基础较弱",
    ]:
        assert line in text
        assert line in captured[0][1]["content"]
    assert text.count("【优化提示】") == 1


def test_fastgpt_wire_prompt_contains_optimization_exactly_once(tmp_path):
    import httpx
    from backend.app.study_guide.llm.client import ModelClient

    store, session, pipeline = setup(tmp_path)
    settings = pipeline.settings
    settings.model_provider = "fastgpt"
    settings.fastgpt_app_id = "test-app"
    settings.fastgpt_url = "https://fastgpt.test/api"
    settings.fastgpt_key = "test-only-key"
    store.update_session(
        session["id"],
        lambda s: s["settings"].update(
            options={"duration": "20 分钟", "focus": ["化学方程式"]}
        ),
    )
    sent = []

    def transport(request):
        body = json.loads(request.content)
        sent.append(body["messages"][0]["content"])
        return httpx.Response(
            200,
            json={
                "choices": [
                    {"message": {"content": json.dumps(SAMPLE, ensure_ascii=False)}}
                ]
            },
        )

    pipeline.model = ModelClient(settings, httpx.MockTransport(transport))
    job = pipeline.start_generation(session["id"], {"prompt": "请生成导学案"})
    pipeline.run(job["id"])
    assert store.get_job(job["id"])["status"] == "completed"
    assert len(sent) == 1
    assert "自学用时：20 分钟" in sent[0]
    assert "本课重点：化学方程式" in sent[0]
    assert sent[0].count("【优化提示】") == 1


def test_safe_blank_header_is_repaired_without_another_model_generation(tmp_path):
    model = Model()
    raw = copy.deepcopy(SAMPLE)
    raw["outline"]["sections"][0]["blocks"][1]["headers"][1] = ""
    calls = []

    def generate(messages):
        calls.append(messages)
        return copy.deepcopy(raw)

    model.generate_json = generate
    store, s, p = setup(tmp_path, model)
    j = p.start_generation(s["id"], {"prompt": "生成"})
    p.run(j["id"])
    done = store.get_job(j["id"])
    assert done["status"] == "completed", done
    v = store.get_session(s["id"])["versions"][-1]
    assert v["content"]["sections"][0]["blocks"][1]["headers"][1] == "填写内容"
    assert v["report"]["format_repairs"][0]["method"] == "automatic"
    assert len(calls) == 1
    from docx import Document

    word = Document(Path(v["directory"]) / "outline.docx")
    assert word.tables[0].cell(0, 1).text == "填写内容"
    assert word.tables[0].cell(1, 1).text == ""


def test_ambiguous_header_uses_local_repair_and_preserves_other_content(tmp_path):
    raw = copy.deepcopy(SAMPLE)
    table = raw["outline"]["sections"][0]["blocks"][1]
    table["headers"][0] = ""
    calls = []

    def generate(messages):
        calls.append(copy.deepcopy(messages))
        if len(calls) == 1:
            return copy.deepcopy(raw)
        # Only this table is supplied, not the full outline or source evidence.
        request = json.loads(messages[-1]["content"])
        assert request["table"]["rows"] == table["rows"]
        assert "outline" not in request and "sources" not in request
        return {"headers": ["反应", "变化的元素", "化合价变化"]}

    model = Model()
    model.generate_json = generate
    store, s, p = setup(tmp_path, model)
    j = p.start_generation(s["id"], {"prompt": "生成"})
    p.run(j["id"])
    assert store.get_job(j["id"])["status"] == "completed", store.get_job(j["id"])
    v = store.get_session(s["id"])["versions"][-1]
    assert len(calls) == 2
    assert v["content"]["sections"][1:] == raw["outline"]["sections"][1:]
    assert v["content"]["sections"][0]["blocks"][1]["rows"] == table["rows"]
    assert v["report"]["format_repairs"][0]["method"] == "model"


def test_unrepairable_header_has_readable_summary_and_separate_detail(tmp_path):
    raw = copy.deepcopy(SAMPLE)
    raw["outline"]["sections"][0]["blocks"][1]["headers"][0] = ""
    calls = []

    def generate(messages):
        calls.append(messages)
        return (
            copy.deepcopy(raw)
            if len(calls) == 1
            else {"headers": ["", "变化的元素", "化合价变化"]}
        )

    model = Model()
    model.generate_json = generate
    store, s, p = setup(tmp_path, model)
    j = p.start_generation(s["id"], {"prompt": "生成"})
    p.run(j["id"])
    done = store.get_job(j["id"])
    assert done["status"] == "failed" and len(calls) == 3
    assert "表头" in done["error"] and "$.outline" not in done["error"]
    assert "headers[0]" in done["error_detail"]
    assert store.get_session(s["id"])["versions"] == []


def test_row_column_mismatch_is_still_rejected_after_repair_attempts(tmp_path):
    raw = copy.deepcopy(SAMPLE)
    raw["outline"]["sections"][0]["blocks"][1]["rows"][0].pop()
    model = Model()
    model.generate_json = lambda _: copy.deepcopy(raw)
    store, s, p = setup(tmp_path, model)
    j = p.start_generation(s["id"], {"prompt": "生成"})
    p.run(j["id"])
    job = store.get_job(j["id"])
    assert job["status"] == "failed" and "列数不一致" in job["error"]
    assert store.get_session(s["id"])["versions"] == []


def test_patch_revision_repairs_only_new_table_headers_and_keeps_history(tmp_path):
    store, s, p = setup(tmp_path)
    first = p.start_generation(s["id"], {"prompt": "生成"})
    p.run(first["id"])
    base = copy.deepcopy(store.get_session(s["id"])["versions"][-1])
    table = copy.deepcopy(base["content"]["sections"][0]["blocks"][1])
    table["headers"][0] = ""
    table["rows"].append(["[Zn] + 2[HCl] = [ZnCl2] + [H2]", None, None])
    calls = []

    def generate(messages):
        calls.append(copy.deepcopy(messages))
        if len(calls) == 1:
            return {
                "mode": "patch",
                "changes": [{"block_id": "table-1", "replacement": [table]}],
            }
        assert json.loads(messages[-1]["content"])["table"]["rows"] == table["rows"]
        return {"headers": ["反应", "变化的元素", "化合价变化"]}

    p.model.generate_json = generate
    job = p.start_revision(
        s["id"], {"prompt": "增加表格中的反应实例", "base_version_id": base["id"]}
    )
    p.run(job["id"])
    assert store.get_job(job["id"])["status"] == "completed", store.get_job(job["id"])
    versions = store.get_session(s["id"])["versions"]
    assert versions[0] == base and len(calls) == 2
    latest = versions[-1]
    assert latest["content"]["sections"][0]["blocks"][1]["rows"] == table["rows"]
    assert latest["content"]["sections"][1:] == base["content"]["sections"][1:]
    saved = json.loads(
        (Path(latest["directory"]) / "generation_report.json").read_text()
    )
    assert saved["format_repairs"] == latest["report"]["format_repairs"]


def test_legacy_equation_is_normalized_in_preview_and_native_word(tmp_path):
    from docx import Document

    model = Model()
    raw = copy.deepcopy(SAMPLE)
    raw["outline"]["sections"][0]["blocks"][0]["content"] = (
        "<br>CaCO3 ===高温=== CaO + CO2↑"
    )
    model.generate_json = lambda _: copy.deepcopy(raw)
    store, session, pipeline = setup(tmp_path, model)
    job = pipeline.start_generation(session["id"], {"prompt": "生成"})
    pipeline.run(job["id"])
    assert store.get_job(job["id"])["status"] == "completed"
    version = store.get_session(session["id"])["versions"][-1]
    runs = version["content"]["sections"][0]["blocks"][0]["content"]
    assert isinstance(runs, list)
    assert {"text": "高温", "condition": True} in runs
    document = Document(Path(version["directory"], "outline.docx"))
    xml = document._element.xml
    assert "===高温===" not in xml and "&lt;br&gt;" not in xml
    assert 'w:val="subscript"' in xml
    assert "w:ruby" in xml
    assert "高温" in xml
    assert any(
        n["method"] == "chemistry_format" for n in version["report"]["format_repairs"]
    )


PROMPT_LEAK = "本导学案严格遵循“通过……能够……”句式设定目标，紧扣“强化化学方程式”的优化提示。所有化学式已按规范添加方括号标记，反应条件采用指定condition对象，表格行内留白使用null。"


def test_prompt_compliance_comment_is_repaired_before_publishing(tmp_path):
    bad = copy.deepcopy(SAMPLE)
    bad["outline"]["thinking"] = [PROMPT_LEAK]
    replies = iter([bad, copy.deepcopy(SAMPLE)])
    messages_seen = []
    model = Model()

    def generate(messages):
        messages_seen.append(copy.deepcopy(messages))
        return next(replies)

    model.generate_json = generate
    store, session, pipeline = setup(tmp_path, model)
    job = pipeline.start_generation(session["id"], {"prompt": "强化化学方程式"})
    pipeline.run(job["id"])
    assert store.get_job(job["id"])["status"] == "completed"
    assert len(messages_seen) == 2
    assert "thinking[0]" in messages_seen[1][-1]["content"]
    version = store.get_session(session["id"])["versions"][-1]
    assert version["content"]["thinking"] == [
        "判断氧化还原反应时，需要逐一比较哪些信息？"
    ]
    assert "condition对象" not in Path(version["directory"], "content.json").read_text()


def test_repeated_prompt_leak_is_not_published(tmp_path):
    bad = copy.deepcopy(SAMPLE)
    bad["outline"]["sections"][0]["blocks"][0]["content"] = PROMPT_LEAK
    model = Model()
    model.generate_json = lambda _: copy.deepcopy(bad)
    store, session, pipeline = setup(tmp_path, model)
    job = pipeline.start_generation(session["id"], {"prompt": "生成"})
    pipeline.run(job["id"])
    done = store.get_job(job["id"])
    assert done["status"] == "failed"
    assert "生成说明" in done["error"]
    assert not store.get_session(session["id"])["versions"]


def test_heating_equation_native_word_preserves_baseline_coefficient(tmp_path):
    from docx import Document

    model = Model()
    raw = copy.deepcopy(SAMPLE)
    raw["outline"]["sections"][0]["blocks"][0]["content"] = (
        "[2KMnO4] △→ [K2MnO4] + [MnO2] + [O2]↑"
    )
    model.generate_json = lambda _: copy.deepcopy(raw)
    store, session, pipeline = setup(tmp_path, model)
    job = pipeline.start_generation(session["id"], {"prompt": "生成"})
    pipeline.run(job["id"])
    assert store.get_job(job["id"])["status"] == "completed"
    version = store.get_session(session["id"])["versions"][-1]
    doc = Document(Path(version["directory"], "outline.docx"))
    para = next(p for p in doc.paragraphs if "KMnO" in p.text)
    xml = para._p.xml
    assert "→" not in xml
    assert "△" in xml and "w:ruby" in xml
    coefficient = next(r for r in para.runs if r.text)
    assert coefficient.text == "2"
    assert not coefficient.font.subscript
    assert any(r.text == "4" and r.font.subscript for r in para.runs)
