import json
from pathlib import Path
import pytest
from docx import Document
from backend.app.study_guide.tools.documents import DocumentEngine

SAMPLE = json.loads(Path("vendor/chemistry/examples/outline.json").read_text())


def test_sources_do_not_gate_rendering_or_add_review_title(tmp_path):
    engine = DocumentEngine()
    assert engine.validate(SAMPLE)["outline"]["sections"]
    report = engine.render(SAMPLE, tmp_path / "v001", "v001", False)
    doc = Document(tmp_path / "v001/outline.docx")
    assert report["draft"]
    assert "待核验" not in doc.paragraphs[0].text
    assert len(doc.tables) == 4
    assert len(doc.tables[-1].rows) == 25
    assert doc.tables[0].cell(1, 1).text == ""
    with pytest.raises(FileExistsError):
        engine.render(SAMPLE, tmp_path / "v001", "v001", True)


def test_revision_hash_and_preservation(tmp_path):
    engine = DocumentEngine()
    report = engine.render(SAMPLE, tmp_path / "v001", "v001", True)
    original = json.loads((tmp_path / "v001/content.json").read_text())
    patch = {
        "base_sha256": report["content_sha256"],
        "changes": [
            {
                "block_id": "question-1",
                "replacement": [{"type": "paragraph", "content": "新问题"}],
            }
        ],
    }
    engine.revise(tmp_path / "v001", patch, tmp_path / "v002", "v002")
    new = json.loads((tmp_path / "v002/content.json").read_text())
    assert new["sections"][1:] == original["sections"][1:]
    patch["base_sha256"] = "stale"
    with pytest.raises(ValueError):
        engine.revise(tmp_path / "v002", patch, tmp_path / "v003", "v003")


def test_legacy_revision_ignores_source_review_gate_without_rewriting_history(tmp_path):
    engine = DocumentEngine()
    report = engine.render(SAMPLE, tmp_path / "v001", "v001", True)
    report_path = tmp_path / "v001/generation_report.json"
    legacy = json.loads(report_path.read_text())
    legacy.update(draft=False, status="generated_pending_visual_review")
    report_path.write_text(json.dumps(legacy, ensure_ascii=False))
    before = report_path.read_bytes()
    patch = {
        "base_sha256": report["content_sha256"],
        "changes": [
            {
                "block_id": "question-1",
                "replacement": [{"type": "paragraph", "content": "新问题"}],
            }
        ],
    }
    revised = engine.revise(tmp_path / "v001", patch, tmp_path / "v002", "v002")
    assert report_path.read_bytes() == before
    assert revised["previous"]["directory"] == str((tmp_path / "v001").resolve())
    assert "待核验" not in Document(tmp_path / "v002/outline.docx").paragraphs[0].text
    original = json.loads((tmp_path / "v001/content.json").read_text())
    updated = json.loads((tmp_path / "v002/content.json").read_text())
    assert original["sections"][1:] == updated["sections"][1:]


def test_rejoined_condition_equation_has_no_internal_word_line_breaks(tmp_path):
    import copy
    from zipfile import ZipFile
    from lxml import etree
    from backend.app.shared.chemistry.formatting import normalize_chemistry

    raw = copy.deepcopy(SAMPLE)
    raw["outline"]["sections"][0]["blocks"].insert(
        0,
        {
            "id": "equation-line",
            "type": "paragraph",
            "content": "C + 2[CuO]\n高温\n= 2[Cu] + [CO2]↑",
        },
    )
    normalized, _ = normalize_chemistry(raw)
    DocumentEngine().render(normalized, tmp_path / "joined", "v001", True)
    ns = {
        "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
        "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
    }
    with ZipFile(tmp_path / "joined/outline.docx") as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
        paragraphs = root.xpath(
            '//w:p[.//m:limUpp/m:lim//m:t[text()="高温"]]', namespaces=ns
        )
        assert len(paragraphs) == 1
        paragraph = paragraphs[0]
        text = "".join(paragraph.xpath(".//w:t/text()", namespaces=ns))
        assert "CuO" in text and "Cu" in text and "CO" in text
        assert not paragraph.xpath(".//w:br", namespaces=ns)
