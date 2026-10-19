import zipfile
from pathlib import Path
from docx import Document
from pypdf import PdfReader

MAX_BYTES = 20 * 1024 * 1024
MAX_TEXT = 100000


def extract_resource(path: Path, original_name: str) -> dict:
    if path.stat().st_size > MAX_BYTES:
        raise ValueError("文件超过20MB，请缩减后重试")
    suffix = Path(original_name).suffix.lower()
    try:
        if suffix in (".txt", ".md"):
            text = path.read_text(encoding="utf-8-sig")
            if "\x00" in text:
                raise ValueError("文件不是有效文本")
            pages = []
        elif suffix == ".docx":
            with zipfile.ZipFile(path) as package:
                if sum(i.file_size for i in package.infolist()) > 50 * 1024 * 1024:
                    raise ValueError("DOCX 解压体积过大")
            doc = Document(path)
            text = "\n".join(
                [p.text for p in doc.paragraphs]
                + [
                    " | ".join(c.text for c in row.cells)
                    for t in doc.tables
                    for row in t.rows
                ]
            )
            pages = []
        elif suffix == ".pdf":
            reader = PdfReader(path)
            if reader.is_encrypted:
                raise ValueError("请上传未加密 PDF")
            if len(reader.pages) > 300:
                raise ValueError("PDF 超过300页，请拆分相关章节")
            pages = [
                {"pdf_page": i + 1, "text": p.extract_text() or ""}
                for i, p in enumerate(reader.pages)
            ]
            text = "\n".join(
                f"[PDF文件页序 {p['pdf_page']}]\n{p['text']}"
                for p in pages
                if p["text"].strip()
            )
            if not text.strip():
                raise ValueError("PDF 没有可提取文字，请先 OCR 或上传文字资料")
        else:
            raise ValueError("支持 TXT、Markdown、DOCX 和文字 PDF")
        if not text.strip():
            raise ValueError("文件内容为空")
        if len(text) > MAX_TEXT:
            raise ValueError("资料超过100000字符，请缩减为本课相关章节")
        return {"text": text, "pages": pages}
    except ValueError:
        raise
    except Exception:
        raise ValueError("文件无法解析，请检查实际格式与扩展名是否一致") from None
